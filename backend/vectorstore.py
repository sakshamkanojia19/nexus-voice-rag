"""Persistent document store backed by Qdrant Cloud + LangChain.

Replaces the old in-memory NumPy store. Documents are parsed, chunked
(LangChain RecursiveCharacterTextSplitter), embedded (OpenAIEmbeddings) and
upserted into a Qdrant collection — so they survive restarts and scale.

Public API: add / search / list / delete / doc_names / has_documents.
IO/embedding calls run in threads to stay async.

`add` and `search` both require an `api_key` because both embed, and embedding
is the only paid call this module makes. Passing it per call rather than
holding one at startup is what lets a public deployment bill each user for
their own documents; see `DocumentStore._vs_for`. The other four methods read
Qdrant payloads directly and cost nothing, so they take no key.
"""
from __future__ import annotations

import asyncio 
import hashlib
import io
import logging
import re
import uuid
from collections import OrderedDict

from langchain_core.documents import Document
from langchain_openai import OpenAIEmbeddings
from langchain_qdrant import QdrantVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client import QdrantClient
from qdrant_client import models as qm

from . import config, providers

log = logging.getLogger("universalcb.vectorstore")

_CHARS_PER_TOKEN = 4
_TARGET_CHARS = config.CHUNK_TOKENS * _CHARS_PER_TOKEN
_OVERLAP_CHARS = config.CHUNK_OVERLAP_TOKENS * _CHARS_PER_TOKEN

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=_TARGET_CHARS,
    chunk_overlap=_OVERLAP_CHARS,
    separators=["\n\n", "\n", ". ", " ", ""],
)


# ── Parsing (PDF / DOCX / text) ──────────────────────────────────────────────
def _ext(filename: str) -> str:
    i = filename.rfind(".")
    return filename[i:].lower() if i >= 0 else ""


def _parse_pdf(data: bytes) -> list[tuple[int | None, str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    out: list[tuple[int | None, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as e:  # noqa: BLE001
            log.debug("pdf page %d extract failed: %s", i, e)
            text = ""
        if text:
            out.append((i, text))
    return out


def _parse_docx(data: bytes) -> list[tuple[int | None, str]]:
    import docx

    document = docx.Document(io.BytesIO(data))
    text = "\n\n".join(p.text.strip() for p in document.paragraphs if p.text and p.text.strip())
    return [(None, text)] if text.strip() else []


def _parse_text(data: bytes) -> list[tuple[int | None, str]]:
    text = data.decode("utf-8", errors="replace").strip()
    return [(None, text)] if text else []


def parse(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    ext = _ext(filename)
    if ext == ".pdf":
        return _parse_pdf(data)
    if ext == ".docx":
        return _parse_docx(data)
    if ext in (".txt", ".md", ".markdown", ".text", ".log", ".csv"):
        return _parse_text(data)
    raise ValueError(f"Unsupported file type '{ext or '(none)'}'. Supported: PDF, DOCX, TXT, MD.")


_WS = re.compile(r"[ \t]+")


def _build_documents(session_id: str, doc_id: str, name: str, segments: list[tuple[int | None, str]]) -> list[Document]:
    docs: list[Document] = []
    for page, text in segments:
        for chunk in _splitter.split_text(text):
            chunk = _WS.sub(" ", chunk).strip()
            if not chunk:
                continue
            docs.append(
                Document(
                    page_content=chunk,
                    metadata={"session_id": session_id, "doc_id": doc_id, "doc_name": name, "page": page},
                )
            )
    return docs


# ── Store ────────────────────────────────────────────────────────────────────
class DocumentStore:
    """Qdrant-backed, per-session document store."""

    def __init__(self) -> None:
        if not config.QDRANT_URL or not config.QDRANT_API_KEY:
            raise RuntimeError("QDRANT_URL / QDRANT_API_KEY not configured in backend/.env")
        # No OPENAI_API_KEY check here on purpose. Embedding is billed to
        # whoever asked for it — see `_vs_for()` — so a public BYOK deployment
        # runs with no server key at all, and refusing to start without one
        # would make that configuration impossible.

        self.collection = config.QDRANT_COLLECTION
        # Filenames per session. Only changes on add/delete, but `_prepare()` needs
        # it on every chat turn — without this we'd scroll the whole collection
        # before each answer. Invalidated in add()/delete().
        self._names_cache: dict[str, list[str]] = {}
        # One vector store per distinct API key, keyed by digest rather than the
        # key itself: dict keys surface in reprs, tracebacks and heap dumps, and
        # a key that reaches a log is a key you have to rotate.
        self._vs_cache: OrderedDict[str, QdrantVectorStore] = OrderedDict()
        self._client = QdrantClient(url=config.QDRANT_URL, api_key=config.QDRANT_API_KEY, timeout=30.0)
        self._ensure_collection()
        self._verify_dim()
        log.info("Qdrant connected: collection=%s", self.collection)

    #: How many per-key vector stores to keep. Each is a thin wrapper around the
    #: shared QdrantClient plus an httpx pool to OpenAI, so the cost of holding
    #: one is small — but unbounded growth on a public deployment is not.
    _VS_CACHE_MAX = 32

    def _verify_dim(self) -> None:
        """Check the collection's vector width against the configured model.

        langchain-qdrant would do this for us, but it does it by *embedding a
        probe string* — one paid call per vector store constructed, which under
        BYOK means charging every user a toll on their first search. Qdrant
        already knows the width, so we ask it instead and pay nobody.
        """
        try:
            info = self._client.get_collection(self.collection)
            params = info.config.params.vectors
            size = params.size if hasattr(params, "size") else params[""].size
        except Exception as e:  # noqa: BLE001
            log.debug("could not read collection vector size: %s", e)
            return
        if size != config.EMBEDDING_DIM:
            raise RuntimeError(
                f"Collection {self.collection!r} stores {size}-dim vectors but "
                f"EMBEDDING_MODEL={config.EMBEDDING_MODEL} produces "
                f"{config.EMBEDDING_DIM}. Point QDRANT_COLLECTION at a new name "
                f"or restore the original model — searching across two embedding "
                f"spaces returns confident nonsense."
            )

    def _vs_for(self, api_key: str) -> QdrantVectorStore:
        """A vector store that embeds on `api_key`'s dime.

        Embedding is the one OpenAI call the document path makes, so whoever's
        key lands here is who pays for uploads and document searches. Binding it
        per request rather than once at startup is what makes BYOK actually
        BYOK: with a server key hard-wired here, `REQUIRE_USER_API_KEY=true`
        would still leave the operator paying for every stranger's uploads.

        Vectors themselves are account-independent — the same model produces the
        same embedding space for everyone — so documents indexed under one key
        stay searchable under another.
        """
        if not api_key:
            raise RuntimeError("No API key available for embedding")
        # The routes refuse non-OpenAI keys before reaching here; this is the
        # backstop that keeps an xAI or Groq key from ever being sent to OpenAI.
        owner = providers.for_key(api_key)
        if owner is not None and not owner.embeddings:
            raise RuntimeError(f"{owner.label} keys cannot embed documents")
        digest = hashlib.sha256(api_key.encode()).hexdigest()
        cached = self._vs_cache.get(digest)
        if cached is not None:
            self._vs_cache.move_to_end(digest)
            return cached

        vs = QdrantVectorStore(
            client=self._client,
            collection_name=self.collection,
            embedding=OpenAIEmbeddings(model=config.EMBEDDING_MODEL, api_key=api_key),
            # Already checked once, for free, in `_verify_dim()`.
            validate_embeddings=False,
            validate_collection_config=False,
        )
        self._vs_cache[digest] = vs
        if len(self._vs_cache) > self._VS_CACHE_MAX:
            self._vs_cache.popitem(last=False)
        return vs

    def _ensure_collection(self) -> None:
        if not self._client.collection_exists(self.collection):
            self._client.create_collection(
                self.collection,
                vectors_config=qm.VectorParams(size=config.EMBEDDING_DIM, distance=qm.Distance.COSINE),
            )
            log.info("Created Qdrant collection %s (dim=%d)", self.collection, config.EMBEDDING_DIM)
        # Qdrant requires payload indexes to filter on these keyword fields.
        # Idempotent: creating an existing index is a no-op / harmless error.
        for field in ("metadata.session_id", "metadata.doc_id"):
            try:
                self._client.create_payload_index(
                    self.collection, field_name=field, field_schema=qm.PayloadSchemaType.KEYWORD,
                )
            except Exception as e:  # noqa: BLE001
                log.debug("payload index %s: %s", field, e)

    def ping(self) -> bool:
        try:
            self._client.get_collections()
            return True
        except Exception:  # noqa: BLE001
            return False

    # -- filters --------------------------------------------------------------
    @staticmethod
    def _filter(session_id: str, doc_ids: list[str] | None = None) -> qm.Filter:
        must: list[qm.FieldCondition] = [
            qm.FieldCondition(key="metadata.session_id", match=qm.MatchValue(value=session_id))
        ]
        if doc_ids is not None:
            must.append(qm.FieldCondition(key="metadata.doc_id", match=qm.MatchAny(any=doc_ids)))
        return qm.Filter(must=must)

    # -- ingest ---------------------------------------------------------------
    async def add(self, session_id: str, filename: str, data: bytes, api_key: str) -> dict:
        doc_id = uuid.uuid4().hex[:12]
        meta = {
            "doc_id": doc_id,
            "name": filename,
            "size": len(data),
            "status": "processing",
            "error": None,
            "n_chunks": 0,
            "pages": [],
        }
        try:
            segments = await asyncio.to_thread(parse, filename, data)
            if not segments:
                raise ValueError("No extractable text found (scanned image PDF?).")
            docs = _build_documents(session_id, doc_id, filename, segments)
            if not docs:
                raise ValueError("Document produced no text chunks.")
            await asyncio.to_thread(self._vs_for(api_key).add_documents, docs)
            meta["status"] = "ready"
            meta["n_chunks"] = len(docs)
            meta["pages"] = sorted({d.metadata["page"] for d in docs if d.metadata.get("page") is not None})
            log.info("ingested %r (%d chunks) for %s", filename, len(docs), session_id)
        except Exception as e:  # noqa: BLE001
            meta["status"] = "error"
            meta["error"] = str(e)
            log.warning("ingest failed for %r: %s", filename, e)
        self._names_cache.pop(session_id, None)
        return meta

    # -- query ----------------------------------------------------------------
    async def search(
        self,
        session_id: str,
        query: str,
        api_key: str,
        k: int | None = None,
        doc_ids: list[str] | None = None,
    ) -> list[dict]:
        if doc_ids is not None and len(doc_ids) == 0:
            return []
        k = k or config.RAG_TOP_K
        flt = self._filter(session_id, doc_ids)
        try:
            results = await asyncio.to_thread(
                self._vs_for(api_key).similarity_search_with_score, query, k, flt
            )
        except Exception as e:  # noqa: BLE001
            log.warning("qdrant search failed: %s", e)
            return []
        out: list[dict] = []
        for doc, score in results:
            m = doc.metadata or {}
            out.append(
                {
                    "doc_id": m.get("doc_id"),
                    "doc_name": m.get("doc_name"),
                    "page": m.get("page"),
                    "text": doc.page_content,
                    "score": float(score),
                }
            )
        return out

    async def list(self, session_id: str) -> list[dict]:
        def _scroll() -> list[dict]:
            docs: dict[str, dict] = {}
            offset = None
            while True:
                points, offset = self._client.scroll(
                    collection_name=self.collection,
                    scroll_filter=self._filter(session_id),
                    with_payload=True,
                    with_vectors=False,
                    limit=256,
                    offset=offset,
                )
                for p in points:
                    m = (p.payload or {}).get("metadata", {})
                    did = m.get("doc_id")
                    if not did:
                        continue
                    entry = docs.setdefault(
                        did,
                        {"doc_id": did, "name": m.get("doc_name", "document"), "size": 0,
                         "status": "ready", "error": None, "n_chunks": 0, "pages": set()},
                    )
                    entry["n_chunks"] += 1
                    if m.get("page") is not None:
                        entry["pages"].add(m["page"])
                if offset is None:
                    break
            for e in docs.values():
                e["pages"] = sorted(e["pages"])
            return list(docs.values())

        try:
            return await asyncio.to_thread(_scroll)
        except Exception as e:  # noqa: BLE001
            log.warning("qdrant list failed: %s", e)
            return []

    async def delete(self, session_id: str, doc_id: str) -> bool:
        def _del() -> bool:
            self._client.delete(
                collection_name=self.collection,
                points_selector=qm.FilterSelector(filter=self._filter(session_id, [doc_id])),
            )
            return True

        try:
            return await asyncio.to_thread(_del)
        except Exception as e:  # noqa: BLE001
            log.warning("qdrant delete failed: %s", e)
            return False
        finally:
            self._names_cache.pop(session_id, None)

    async def doc_names(self, session_id: str) -> list[str]:
        """Cached — called on every chat turn, so it must not hit the network."""
        cached = self._names_cache.get(session_id)
        if cached is not None:
            return cached
        names = [d["name"] for d in await self.list(session_id)]
        self._names_cache[session_id] = names
        return names

    async def has_documents(self, session_id: str) -> bool:
        """Cheap existence check — fetch at most one point instead of scrolling
        every chunk in the session (which list() does)."""
        if self._names_cache.get(session_id):
            return True

        def _probe() -> bool:
            points, _ = self._client.scroll(
                collection_name=self.collection,
                scroll_filter=self._filter(session_id),
                with_payload=False,
                with_vectors=False,
                limit=1,
            )
            return len(points) > 0

        try:
            return await asyncio.to_thread(_probe)
        except Exception as e:  # noqa: BLE001
            log.warning("qdrant existence probe failed: %s", e)
            return False
